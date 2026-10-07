# -*- coding: utf-8 -*-
"""Extrait les fichiers sources (Excel/Word) vers data/source/*.json.

Usage :
    python extract_sources.py <dossier des sources>

Le dossier doit contenir Russe.xlsx, italien vocabulaire.ods.xlsx et
Italien.docx (grammaire, extraite à part).
"""
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

DESKTOP = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
OUT = Path(__file__).resolve().parent.parent / "data" / "source"
OUT.mkdir(parents=True, exist_ok=True)


def dump(name, obj):
    path = OUT / name
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{name}: {len(obj)} items")


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def extract_russe():
    wb = openpyxl.load_workbook(DESKTOP / "Russe.xlsx", read_only=True, data_only=True)

    # Voc : Russe | Traduction | (col C libre) | cacher | poids | ID | Prononciation
    voc = []
    rows = list(wb["Voc"].iter_rows(values_only=True))
    for r in rows[1:]:  # skip header
        r = list(r) + [None] * 10
        ru, fr = clean(r[0]), clean(r[1])
        if not ru or not fr:
            continue
        # la colonne prononciation n'est pas toujours au même index : cherche
        # une chaîne latine après les 2 premières colonnes
        pron = None
        for c in r[2:10]:
            c = clean(c)
            if c and len(c) > 1 and not re.search(r"[а-яА-ЯёЁ]", c) and not re.fullmatch(r"[ynl]|[\d.,]+", c):
                pron = c
                break
        voc.append({"ru": ru, "fr": fr, "pron": pron})
    dump("ru_vocab.json", voc)

    phrases = []
    rows = list(wb["Phrases"].iter_rows(values_only=True))
    for r in rows[1:]:
        r = list(r) + [None] * 10
        ru, fr = clean(r[0]), clean(r[1])
        if not ru or not fr:
            continue
        pron = None
        for c in r[2:10]:
            c = clean(c)
            if c and len(c) > 2 and not re.search(r"[а-яА-ЯёЁ]", c) and not re.fullmatch(r"[ynl]|[\d.,]+", c):
                pron = c
                break
        phrases.append({"ru": ru, "fr": fr, "pron": pron})
    dump("ru_phrases.json", phrases)

    chiffres = []
    for r in wb["Chiffres"].iter_rows(values_only=True):
        r = list(r) + [None] * 4
        if clean(r[0]) and clean(r[1]):
            chiffres.append({"ru": clean(r[0]), "fr": clean(r[1]),
                             "ordinal_ru": clean(r[2]), "ordinal_fr": clean(r[3])})
    dump("ru_chiffres.json", chiffres)

    autre = []
    for r in wb["Autre voc"].iter_rows(values_only=True):
        r = list(r) + [None] * 2
        a, b = clean(r[0]), clean(r[1])
        if a and b:
            # sens variable dans cet onglet : russe/translit à gauche ou à droite
            if re.search(r"[а-яА-ЯёЁ]", a):
                autre.append({"ru": a, "fr": b, "pron": None})
            else:
                autre.append({"ru": None, "fr": a, "pron": b})
    dump("ru_autre_voc.json", autre)

    # Onglets de grammaire : texte brut par onglet
    grammaire = {}
    for sheet in ["Verbes", "Grammaire", "Pronoms", "Adjectifs", "Genre", "Pluriel", "Prononciation"]:
        lines = []
        for r in wb[sheet].iter_rows(values_only=True):
            cells = [clean(c) for c in r if clean(c)]
            if cells:
                lines.append(" | ".join(cells))
        grammaire[sheet] = lines
    (OUT / "ru_grammaire_brut.json").write_text(
        json.dumps(grammaire, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"ru_grammaire_brut.json: {sum(len(v) for v in grammaire.values())} lignes")


def extract_italien():
    wb = openpyxl.load_workbook(DESKTOP / "italien vocabulaire.ods.xlsx",
                                read_only=True, data_only=True)
    items = []
    for sheet, statut in [("italien_vocabulaire", "actif"), ("aquis", "ex-aquis")]:
        for r in wb[sheet].iter_rows(values_only=True):
            r = list(r) + [None] * 4
            it, fr = clean(r[1]), clean(r[2])
            if it and fr:
                items.append({"it": it, "fr": fr, "source": statut})
    dump("it_vocab.json", items)


def extract_docx():
    z = zipfile.ZipFile(DESKTOP / "Italien.docx")
    xml = z.read("word/document.xml").decode("utf-8")
    paras = []
    for p in re.split(r"</w:p>", xml):
        t = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p))
        t = t.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&").strip()
        if t:
            paras.append(t)
    dump("it_grammaire_brut.json", paras)


if __name__ == "__main__":
    extract_russe()
    extract_italien()
    extract_docx()
    print("OK ->", OUT)
